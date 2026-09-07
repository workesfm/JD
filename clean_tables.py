#!/usr/bin/env python3
"""Normalize flat CSV/XLSX exports with an explicit, auditable rule set."""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import io
import json
import re
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, localcontext
from pathlib import Path

FIELDS = ["record_id", "date", "amount", "currency", "description"]
ORIGIN = ["source_file", "source_record"]
AMOUNT = re.compile(r"[+-]?\d+(?:\.\d{1,2})?\Z")


def text_value(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return str(value).strip()


def parse_date(value, formats):
    matches = set()
    for fmt in formats:
        try:
            matches.add(datetime.strptime(value, fmt).date().isoformat())
        except ValueError:
            pass
    if len(matches) != 1:
        raise ValueError("ambiguous_date" if len(matches) > 1 else "invalid_date")
    return matches.pop()


def normalize(raw, config):
    row = {field: text_value(raw.get(field)) for field in FIELDS}
    if not row["record_id"]:
        raise ValueError("missing_record_id")
    if row["record_id"].startswith("="):
        raise ValueError("formula_in_record_id")
    row["date"] = parse_date(row["date"], config["date_formats"])
    if not AMOUNT.fullmatch(row["amount"]):
        raise ValueError("invalid_amount_or_more_than_2_decimal_places")
    amount = Decimal(row["amount"])
    if abs(amount) >= Decimal("1000000000000"):
        raise ValueError("amount_outside_supported_range")
    row["amount"] = format((amount if amount else Decimal("0")).quantize(Decimal("0.01")), ".2f")
    row["currency"] = row["currency"].upper()
    if row["currency"] not in config["currencies"]:
        raise ValueError("missing_or_unsupported_currency")
    return row


def read_table(path, config):
    payload = path.read_bytes()
    if len(payload) > 10 * 1024 * 1024:
        raise ValueError(f"Input exceeds 10 MiB: {path.name}")
    manifest = {"filename": path.name, "bytes": len(payload),
                "sha256": hashlib.sha256(payload).hexdigest()}
    workbook = None
    if path.suffix.lower() == ".csv":
        rows = csv.reader(io.StringIO(payload.decode(config.get("encoding", "utf-8-sig")), newline=""), strict=True)
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        sheet_name = config.get("xlsx_sheet")
        if sheet_name is None and len(workbook.sheetnames) != 1:
            workbook.close()
            raise ValueError("Select xlsx_sheet explicitly for a multi-sheet workbook")
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"Worksheet not found: {sheet_name}")
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        manifest["worksheet"] = sheet.title
    else:
        raise ValueError("Only .csv and .xlsx inputs are supported")
    try:
        header = [text_value(x) for x in next(rows, [])]
        if not header or any(not h for h in header) or len(set(header)) != len(header):
            raise ValueError(f"Empty or duplicate column headers: {path.name}")
        indices = {}
        for field in FIELDS:
            aliases = config["columns"].get(field, [field])
            found = [index for index, h in enumerate(header) if h in aliases]
            if len(found) > 1 or (not found and field != "description"):
                raise ValueError(f"Missing or ambiguous column {field}: {path.name}")
            if found:
                indices[field] = found[0]
        if len(set(indices.values())) != len(indices):
            raise ValueError("One input column cannot map to multiple output fields")
        records = []
        for number, values in enumerate(rows, start=2):
            if not any(text_value(x) for x in values):
                continue
            if len(records) >= config.get("max_rows", 50000):
                raise ValueError("Input exceeds configured row limit")
            original = [text_value(x) for x in values]
            raw = {field: original[index] if index < len(original) else "" for field, index in indices.items()}
            record = {"source_file": path.name, "source_record": number, "raw": raw,
                      "original": original, "record_id": raw.get("record_id", "")}
            if len(values) != len(header):
                record["reason"] = "column_count_mismatch"
            elif not isinstance(values[indices["record_id"]], str):
                record["reason"] = "record_id_must_be_stored_as_text"
            records.append(record)
        manifest["header"] = header
        return records, manifest
    finally:
        if workbook is not None:
            workbook.close()


def clean(records, config):
    groups = defaultdict(list)
    exceptions = []
    for record in records:
        candidate = dict(record)
        if "reason" not in candidate:
            try:
                candidate["normalized"] = normalize(candidate["raw"], config)
            except ValueError as error:
                candidate["reason"] = str(error)
        key = text_value(candidate.get("record_id"))
        if not key:
            exceptions.append(candidate)
        else:
            groups[key].append(candidate)

    accepted, duplicates = [], []
    for group in groups.values():
        invalid = any("reason" in row for row in group)
        signatures = {tuple(row["normalized"][f] for f in FIELDS) for row in group if "normalized" in row}
        if invalid or len(signatures) > 1:
            for row in group:
                row.setdefault("reason", "same_id_has_invalid_row" if invalid else "conflicting_record_id")
                exceptions.append(row)
            continue
        first, *extra = group
        accepted.append({**first["normalized"], **{f: first[f] for f in ORIGIN}})
        for row in extra:
            row["reason"] = "exact_duplicate_after_normalization"
            row["retained_source"] = f'{first["source_file"]}:{first["source_record"]}'
            duplicates.append(row)

    totals = defaultdict(lambda: Decimal("0.00"))
    with localcontext() as context:
        context.prec = 50
        for row in accepted:
            totals[row["currency"]] += Decimal(row["amount"])
    summary = {"input_records": len(records), "accepted_records": len(accepted),
               "duplicate_records": len(duplicates), "exception_records": len(exceptions),
               "accepted_totals_by_currency": {k: format(v, ".2f") for k, v in sorted(totals.items())}}
    if sum(summary[k] for k in ("accepted_records", "duplicate_records", "exception_records")) != len(records):
        raise RuntimeError("Row accounting failed")
    return accepted, duplicates, exceptions, summary


def audit_rows(records):
    return [{**{f: row.get(f, "") for f in ORIGIN + ["record_id", "reason", "retained_source"]},
             "original_values_json": json.dumps(row["original"], ensure_ascii=False),
             "mapped_values_json": json.dumps(row["raw"], ensure_ascii=False)} for row in records]


def spreadsheet_safe(value, numeric=False):
    value = str(value)
    if not numeric and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_csv(path, columns, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: spreadsheet_safe(row.get(k, ""), numeric=k in ("amount", "source_record")) for k in columns})


def write_workbook(path, accepted, duplicates, exceptions, summary):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    book.remove(book.active)
    audit_columns = ORIGIN + ["record_id", "reason", "retained_source", "original_values_json", "mapped_values_json"]
    summary_rows = [{"metric": k, "value": v} for k, v in summary.items() if isinstance(v, int)]
    summary_rows.extend({"metric": f"accepted_total_{currency}", "value": amount} for currency, amount in summary["accepted_totals_by_currency"].items())
    for name, columns, rows in [("Clean", FIELDS + ORIGIN, accepted), ("Duplicates", audit_columns, duplicates),
                                ("Exceptions", audit_columns, exceptions), ("Summary", ["metric", "value"], summary_rows)]:
        sheet = book.create_sheet(name)
        sheet.append(columns)
        for row_number, row in enumerate(rows, start=2):
            for column_number, key in enumerate(columns, start=1):
                value = float(row[key]) if key == "amount" else row.get(key, "")
                cell = sheet.cell(row=row_number, column=column_number, value=value)
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Keep formula-looking customer text as inert text.
                if key == "amount":
                    cell.number_format = "0.00"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="183153")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cells in sheet.columns:
            sheet.column_dimensions[cells[0].column_letter].width = min(60, max(16, len(str(cells[0].value)) + 3))
    book.save(path)


def write_report(path, summary, exceptions):
    counts = [("Input records", summary["input_records"]), ("Accepted", summary["accepted_records"]),
              ("Duplicates set aside", summary["duplicate_records"]), ("Needs review", summary["exception_records"])]
    cards = "".join(f'<div class="card"><span>{label}</span><strong>{value}</strong></div>' for label, value in counts)
    totals = "".join(f'<tr><td>{html.escape(currency)}</td><td>{html.escape(amount)}</td></tr>' for currency, amount in summary["accepted_totals_by_currency"].items())
    problems = "".join(f'<tr><td>{html.escape(str(row["source_file"]))}:{row["source_record"]}</td><td>{html.escape(str(row["record_id"]))}</td><td>{html.escape(str(row["reason"]))}</td></tr>' for row in exceptions[:50])
    path.write_text(f'''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Table cleanup — verification report</title><style>
body{{font:16px/1.6 system-ui,sans-serif;background:#f3f6fa;color:#183153;margin:0}}main{{max-width:960px;margin:auto;padding:48px 24px}}
h1{{font-size:36px;line-height:1.2}}.eyebrow{{font-size:12px;letter-spacing:.15em;color:#436180}}.grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:32px 0}}
.card,section{{background:white;padding:24px;border-radius:12px;border:1px solid #dce5ef}}.card span{{display:block;font-size:13px}}strong{{display:block;font-size:36px}}section{{margin:20px 0}}table{{width:100%;border-collapse:collapse}}td,th{{text-align:left;padding:12px 8px;border-bottom:1px solid #e6ecf2}}.note{{color:#547086;font-size:14px}}a{{color:#1567ac}}@media(max-width:650px){{.grid{{grid-template-columns:repeat(2,1fr)}}h1{{font-size:28px}}td{{overflow-wrap:anywhere}}}}
</style><main><div class="eyebrow">DATA CLEANUP · VERIFICATION REPORT</div><h1>Every input record accounted for.</h1>
<p>Accepted records, duplicates and exceptions remain traceable to their source. Confirm exception decisions before using these totals as a complete business report.</p>
<div class="grid">{cards}</div><section><h2>Accepted totals</h2><table><tr><th>Currency</th><th>Exact decimal total</th></tr>{totals}</table><p class="note">Totals include accepted records only. Different currencies are never combined.</p></section>
<section><h2>Exceptions to review</h2><table><tr><th>Source record</th><th>Record ID</th><th>Reason</th></tr>{problems}</table><p class="note">First 50 exceptions shown. The workbook and CSV contain every exception. Both sides of a conflicting ID are held for review.</p></section>
<p><a href="cleaned.xlsx">Excel workbook</a> · <a href="cleaned.csv">Clean CSV</a> · <a href="exceptions.csv">All exceptions</a> · <a href="summary.json">Rules and input hashes</a></p>
</main></html>''', encoding="utf-8")


def run(inputs, config_path, output):
    config_bytes = config_path.read_bytes()
    config = json.loads(config_bytes)
    if not config.get("date_formats") or not config.get("currencies") or not isinstance(config.get("columns"), dict):
        raise ValueError("Config needs date_formats, currencies and columns")
    if not inputs or len(inputs) > 5:
        raise ValueError("Provide between 1 and 5 input files")
    if len({p.name for p in inputs}) != len(inputs):
        raise ValueError("Input basenames must be unique for unambiguous source references")
    all_records, manifest = [], []
    for path in inputs:
        rows, source = read_table(path, config)
        all_records.extend(rows)
        manifest.append(source)
        if len(all_records) > config.get("max_rows", 50000):
            raise ValueError("Combined inputs exceed configured row limit")
    accepted, duplicates, exceptions, summary = clean(all_records, config)
    summary.update({"generated_at_utc": datetime.now(timezone.utc).isoformat(), "inputs": manifest,
                    "config_sha256": hashlib.sha256(config_bytes).hexdigest(), "rules": config,
                    "limits": "Flat tables; no formulas evaluated, currency conversion or OCR. Summary covers accepted rows only."})
    output.mkdir(parents=True, exist_ok=False)
    audit_columns = ORIGIN + ["record_id", "reason", "retained_source", "original_values_json", "mapped_values_json"]
    duplicate_audit, exception_audit = audit_rows(duplicates), audit_rows(exceptions)
    write_csv(output / "cleaned.csv", FIELDS + ORIGIN, accepted)
    write_csv(output / "duplicates.csv", audit_columns, duplicate_audit)
    write_csv(output / "exceptions.csv", audit_columns, exception_audit)
    write_workbook(output / "cleaned.xlsx", accepted, duplicate_audit, exception_audit, summary)
    (output / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    write_report(output / "report.html", summary, exception_audit)
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--config", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path, help="New directory; existing directories are never overwritten")
    args = parser.parse_args()
    try:
        result = run(args.inputs, args.config, args.out)
    except (ValueError, OSError, UnicodeError, csv.Error) as error:
        parser.exit(2, f"Input/output error: {error}\n")
    print(json.dumps({k: result[k] for k in ["input_records", "accepted_records", "duplicate_records", "exception_records", "accepted_totals_by_currency"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
