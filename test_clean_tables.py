import copy
import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from clean_tables import clean, normalize, parse_date, read_table, run

SAMPLE = Path(__file__).parent / "sample"
CONFIG = json.loads((SAMPLE / "config.json").read_text())


def record(identifier="0001", amount="1.00", currency="USD", row=2):
    raw = dict(record_id=identifier, date="2026-09-01", amount=amount, currency=currency, description="Sample")
    return dict(source_file="input.csv", source_record=row, raw=raw, original=list(raw.values()), record_id=identifier)


class CleanupTests(unittest.TestCase):
    def test_sample_accounting_and_money(self):
        rows = []
        for name in ["batch-a.csv", "batch-b.csv"]:
            rows.extend(read_table(SAMPLE / name, CONFIG)[0])
        accepted, duplicates, exceptions, summary = clean(rows, CONFIG)
        self.assertEqual((len(accepted), len(duplicates), len(exceptions)), (5, 1, 6))
        self.assertEqual(summary["input_records"], 12)
        self.assertEqual(summary["accepted_totals_by_currency"], {"CNY": "88.00", "USD": "110.30"})
        self.assertEqual(accepted[0]["record_id"], "0001")
        self.assertEqual({r["raw"]["amount"] for r in exceptions if r["record_id"] == "0002"}, {"25.50", "26.50"})

    def test_conflict_never_chooses_first_or_last(self):
        rows = [record(amount="20"), record(amount="21", row=3)]
        for order in [rows, list(reversed(rows))]:
            accepted, duplicates, exceptions, _ = clean(order, CONFIG)
            self.assertEqual((len(accepted), len(duplicates), len(exceptions)), (0, 0, 2))

    def test_invalid_sibling_quarantines_whole_id(self):
        accepted, _, exceptions, _ = clean([record(), record(amount="bad", row=3)], CONFIG)
        self.assertEqual(accepted, [])
        self.assertEqual({r["reason"] for r in exceptions}, {"same_id_has_invalid_row", "invalid_amount_or_more_than_2_decimal_places"})

    def test_amount_rules_and_exact_decimal_sum(self):
        for amount in ["1.234", "NaN", "Infinity", "$1.00", "1,00", "1e2", "1000000000000"]:
            with self.subTest(amount=amount), self.assertRaises(ValueError):
                normalize(record(amount=amount)["raw"], CONFIG)
        _, _, _, summary = clean([record("a", "0.10"), record("b", "0.20"), record("c", "-0.05"), record("d", "100", "CNY")], CONFIG)
        self.assertEqual(summary["accepted_totals_by_currency"], {"USD": "0.25", "CNY": "100.00"})

    def test_dates_are_not_guessed(self):
        with self.assertRaisesRegex(ValueError, "ambiguous_date"):
            parse_date("03/04/2026", ["%d/%m/%Y", "%m/%d/%Y"])
        with self.assertRaises(ValueError):
            parse_date("2026-02-29", CONFIG["date_formats"])
        self.assertEqual(parse_date("2024-02-29", CONFIG["date_formats"]), "2024-02-29")

    def test_workbook_csv_and_formula_text(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "delivery"
            run([SAMPLE / "batch-a.csv", SAMPLE / "batch-b.csv"], SAMPLE / "config.json", output)
            book = load_workbook(output / "cleaned.xlsx", data_only=False)
            self.assertEqual(book.sheetnames, ["Clean", "Duplicates", "Exceptions", "Summary"])
            rows = list(book["Clean"].iter_rows(min_row=2))
            self.assertEqual(rows[0][0].value, "0001")
            formula_like = next(row[4] for row in rows if row[0].value == "0008")
            self.assertEqual((formula_like.value, formula_like.data_type), ("=1+1", "s"))
            book.close()
            with (output / "cleaned.csv").open(encoding="utf-8-sig", newline="") as stream:
                exported = list(csv.DictReader(stream))
            self.assertEqual(next(row["description"] for row in exported if row["record_id"] == "0008"), "'=1+1")
            self.assertEqual(next(row["amount"] for row in exported if row["record_id"] == "0003"), "-10.00")
            with self.assertRaises(FileExistsError):
                run([SAMPLE / "batch-a.csv"], SAMPLE / "config.json", output)

    def test_xlsx_requires_explicit_sheet_and_text_ids(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Data"
            sheet.append(list(record()["raw"]))
            sheet.append(["0001", "2026-09-01", 1.2, "USD", "safe"])
            sheet.append([2, "2026-09-01", 2, "USD", "numeric ID"])
            sheet["A3"].number_format = "0000"
            sheet.append(["0003", "2026-09-01", "=1+1", "USD", "formula amount"])
            book.create_sheet("Notes")
            book.save(path)
            with self.assertRaisesRegex(ValueError, "Select xlsx_sheet"):
                read_table(path, CONFIG)
            config = {**CONFIG, "xlsx_sheet": "Data"}
            rows, _ = read_table(path, config)
            accepted, _, exceptions, _ = clean(rows, config)
            self.assertEqual([r["record_id"] for r in accepted], ["0001"])
            self.assertEqual(len(exceptions), 2)
            self.assertIn("record_id_must_be_stored_as_text", {r["reason"] for r in exceptions})

    def test_schema_errors_and_extra_columns(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.csv"
            path.write_text("record_id,date,amount,currency\n0001,2026-09-01,1,USD,extra\n")
            rows, _ = read_table(path, CONFIG)
            accepted, _, exceptions, _ = clean(rows, CONFIG)
            self.assertEqual(accepted, [])
            self.assertEqual(exceptions[0]["reason"], "column_count_mismatch")
            path.write_text("record_id,date,amount,amount,currency\n0001,2026-09-01,1,1,USD\n")
            with self.assertRaisesRegex(ValueError, "duplicate column"):
                read_table(path, CONFIG)

    def test_empty_data_and_missing_currency(self):
        _, _, _, summary = clean([], CONFIG)
        self.assertEqual(summary["input_records"], 0)
        self.assertEqual(summary["accepted_totals_by_currency"], {})
        accepted, _, exceptions, _ = clean([record(currency="")], CONFIG)
        self.assertEqual(accepted, [])
        self.assertEqual(exceptions[0]["reason"], "missing_or_unsupported_currency")


if __name__ == "__main__":
    unittest.main()
